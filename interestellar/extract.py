"""Extracción de texto de los archivos que sube el administrador.

Devuelve una lista de páginas: [{"page": n, "text": "..."}]. Para formatos sin
páginas (DOCX, TXT) se devuelve una sola "página" con page=None.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

SUPPORTED_EXTENSIONS = [".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".html", ".htm"]


def is_supported(filename: str) -> bool:
    return Path(filename).suffix.lower() in SUPPORTED_EXTENSIONS


def extract_pages(data: bytes, filename: str) -> list[dict]:
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return [{"page": i + 1, "text": _clean(p.extract_text() or "")} for i, p in enumerate(reader.pages)]

    if ext == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return [{"page": None, "text": _clean("\n".join(parts))}]

    if ext in (".html", ".htm"):
        html = data.decode("utf-8", errors="replace")
        text = re.sub(r"<(script|style)[\s\S]*?</\1>", " ", html, flags=re.I)
        text = re.sub(r"<[^>]+>", " ", text).replace("&nbsp;", " ").replace("&amp;", "&")
        return [{"page": None, "text": _clean(text)}]

    if ext in SUPPORTED_EXTENSIONS:
        return [{"page": None, "text": _clean(data.decode("utf-8", errors="replace"))}]

    raise ValueError(f"Formato no soportado: {ext or 'sin extensión'}")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# Divide cada página en fragmentos de ~CHUNK_SIZE caracteres con solapamiento,
# cortando preferentemente en saltos de párrafo o final de frase.
CHUNK_SIZE = 1400
OVERLAP = 250


def chunk_pages(pages: list[dict]) -> list[dict]:
    chunks = []
    for p in pages:
        text = p["text"]
        if not text:
            continue
        start = 0
        while start < len(text):
            end = min(start + CHUNK_SIZE, len(text))
            if end < len(text):
                window = text[start:end]
                cut = max(window.rfind("\n\n"), window.rfind(". "))
                if cut > CHUNK_SIZE * 0.5:
                    end = start + cut + 1
            piece = text[start:end].strip()
            if piece:
                chunks.append({"page": p["page"], "text": piece})
            if end >= len(text):
                break
            start = end - OVERLAP
    return chunks
